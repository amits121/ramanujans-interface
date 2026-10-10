#!/usr/bin/env python3
"""
Build the Techdocumentation twins (.docx and .txt) from their markdown sources in "md files/".

    python3 scripts/build_techdocs.py "md files/Ramanujans-Interface-Master-ToDo.md" [more .md files]

Format: Letter, 1 inch margins, Arial 12, Normal paragraphs only (no heading styles), bold section
lines, real tables with a shaded header row, checklists as ☑ / ☐ paragraphs, bullets as List Bullet.
Markdown understood: "# title" (first one is the document title), "## section", "### subsection",
"- [x] / - [ ]" checklist items, "- " bullets, "_note_" italic lines, "| a | b |" tables (a |---| row
is skipped), blank-line separated paragraphs. Backticks and ** are dropped.

Needs python-docx (the Mac's framework Python has it). Documents only: nothing here touches the engine.
Copyright (c) 2026 Intelligent Cloud Lab Inc.
"""

import os
import re
import sys

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, 'Techdocumentation')
FONT = 'Arial'


def clean(text):
    return re.sub(r'\*\*(.+?)\*\*', r'\1', text.replace('`', '')).strip()


def new_doc():
    d = Document()
    s = d.sections[0]
    s.page_width, s.page_height = Inches(8.5), Inches(11)
    for side in ('left_margin', 'right_margin', 'top_margin', 'bottom_margin'):
        setattr(s, side, Inches(1))
    for name in ('Normal', 'List Bullet'):
        st = d.styles[name]
        st.font.name, st.font.size = FONT, Pt(12)
        st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        st.paragraph_format.space_after = Pt(6)
        st.paragraph_format.line_spacing = 1.15
    return d


def para(d, text, bold=False, italic=False, before=0, after=6, style=None):
    p = d.add_paragraph(style=style) if style else d.add_paragraph()
    r = p.add_run(text)
    r.bold, r.italic = bold, italic
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    return p


def table(d, rows):
    width = max(len(r) for r in rows)
    rows = [r + [''] * (width - len(r)) for r in rows]
    t = d.add_table(rows=len(rows), cols=width)
    t.style = 'Table Grid'
    for ri, row in enumerate(rows):
        for ci, cell in enumerate(row):
            c = t.cell(ri, ci)
            c.text = ''
            p = c.paragraphs[0]
            r = p.add_run(cell)
            r.font.name, r.font.size, r.bold = FONT, Pt(9), ri == 0
            p.paragraph_format.space_after = Pt(0)
            if ri == 0:
                shd = OxmlElement('w:shd')
                shd.set(qn('w:val'), 'clear')
                shd.set(qn('w:color'), 'auto')
                shd.set(qn('w:fill'), 'EEEEEE')
                c._tc.get_or_add_tcPr().append(shd)
    d.add_paragraph().paragraph_format.space_after = Pt(2)


def split_row(line):
    return [clean(c) for c in line.strip().strip('|').split('|')]


def parse(md_text):
    """Yield (kind, payload) blocks: title, section, sub, note, check, bullet, table, para."""
    lines = md_text.split('\n')
    buf, tbl, seen_title = [], [], False

    def flush():
        nonlocal buf, tbl
        if tbl:
            yield_blocks.append(('table', tbl))
            tbl = []
        if buf:
            yield_blocks.append(('para', clean(' '.join(s.strip() for s in buf))))
            buf = []

    yield_blocks = []
    for raw in lines:
        s = raw.rstrip()
        if s.strip().startswith('|'):
            if buf:
                flush()
            if re.match(r'^\s*\|?\s*:?-{2,}', s):
                continue
            tbl.append(split_row(s))
            continue
        if tbl:
            flush()
        if not s.strip():
            flush()
            continue
        if s.startswith('# '):
            flush()
            yield_blocks.append(('section' if seen_title else 'title', clean(s[2:])))
            seen_title = True
        elif s.startswith('## '):
            flush()
            yield_blocks.append(('section', clean(s[3:])))
        elif s.startswith('### '):
            flush()
            yield_blocks.append(('sub', clean(s[4:])))
        elif re.match(r'^- \[[xX ]\] ', s):
            flush()
            yield_blocks.append(('check', ('☑ ' if s[3].lower() == 'x' else '☐ ') + clean(s[6:])))
        elif s.startswith('- '):
            flush()
            yield_blocks.append(('bullet', clean(s[2:])))
        elif re.match(r'^_.*_$', s.strip()):
            flush()
            yield_blocks.append(('note', clean(s.strip()[1:-1])))
        else:
            buf.append(s)
    flush()
    return yield_blocks


def build(md_path):
    blocks = parse(open(md_path, encoding='utf-8').read())
    d = new_doc()
    txt = []
    for kind, payload in blocks:
        if kind == 'title':
            para(d, payload, bold=True, after=8)
            txt.append(payload)
        elif kind == 'section':
            para(d, payload, bold=True, before=14, after=6)
            txt += ['', payload]
        elif kind == 'sub':
            para(d, payload, bold=True, before=10, after=4)
            txt += ['', payload]
        elif kind == 'note':
            para(d, payload, italic=True)
            txt.append(payload)
        elif kind == 'check':
            para(d, payload, after=2)
            txt.append(payload)
        elif kind == 'bullet':
            para(d, payload, after=3, style='List Bullet')
            txt.append('- ' + payload)
        elif kind == 'table':
            table(d, payload)
            for row in payload:
                txt.extend(row)
        else:
            para(d, payload)
            txt.append(payload)
    name = os.path.splitext(os.path.basename(md_path))[0]
    docx_path = os.path.join(OUT_DIR, name + '.docx')
    txt_path = os.path.join(OUT_DIR, name + '.txt')
    d.save(docx_path)
    with open(txt_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(txt).strip() + '\n')
    print(f'{name}: {len(blocks)} blocks -> {os.path.relpath(docx_path, ROOT)}, {os.path.relpath(txt_path, ROOT)}')


if __name__ == '__main__':
    if len(sys.argv) < 2:
        sys.exit(__doc__)
    for path in sys.argv[1:]:
        build(path)
